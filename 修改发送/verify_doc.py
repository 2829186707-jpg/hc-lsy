#!/usr/bin/env python
# -*- coding: utf-8 -*-
from docx import Document

doc = Document(r'C:\Users\veken\Desktop\hc-lsy\修改发送\情侣空间部署指南.docx')
print('=== 文档结构验证 ===')
print(f'总段落数: {len(doc.paragraphs)}')
print()
headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
print(f'标题数: {len(headings)}')
for h in headings:
    level = h.style.name.replace('Heading ', '')
    indent = '  ' * (int(level) - 1) if level.isdigit() else ''
    print(f'{indent}[H{level}] {h.text}')
print()
full_text = '\n'.join([p.text for p in doc.paragraphs])
checks = [
    ('GitHub 注册步骤', 'github.com' in full_text.lower()),
    ('私有仓库说明', 'Private' in full_text),
    ('上传文件说明', 'uploading an existing file' in full_text),
    ('GitHub Pages 配置', 'Deploy from a branch' in full_text),
    ('默认密码', '123456' in full_text),
    ('Token 生成步骤', 'Fine-grained tokens' in full_text),
    ('Token 只显示一次警告', '只显示这一次' in full_text),
    ('iPhone 添加桌面', 'Safari' in full_text),
    ('Android 添加桌面', 'Chrome' in full_text),
    ('修改显示名字', '真实名字' in full_text),
    ('常见问题', '常见问题' in full_text),
    ('功能一览', '功能一览' in full_text),
    ('数据备份提醒', '导出数据' in full_text),
    ('隐私提醒', '私有' in full_text),
    ('.gitignore 隐藏文件提醒', '隐藏文件' in full_text),
    ('不勾选 README 警告', '不要勾选' in full_text),
]
print('=== 内容检查 ===')
all_pass = True
for name, result in checks:
    status = 'PASS' if result else 'FAIL'
    if not result:
        all_pass = False
    print(f'  [{status}] {name}')
print()
result_text = '全部通过' if all_pass else '存在未通过项'
print(f'总体结果: {result_text}')
